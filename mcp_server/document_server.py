"""
mcp_server/document_server.py — MCP Document Server
=====================================================

PURPOSE & DESIGN:
    This module implements the Model Context Protocol (MCP) server for
    LegalShield AI. It acts as a secure, local intermediary between the AI
    agents and the file system / legal reference library.

    The server is launched as a subprocess (stdio transport) by the Google
    Antigravity SDK when an agent session starts. All communication happens
    over standard input/output — no network ports are opened, maintaining
    a strong security posture.

WHY MCP?
    The MCP architecture separates concerns clearly:
    - Agents (AI reasoning) are responsible for THINKING about the content.
    - The MCP server is responsible for FETCHING the content.
    This means agents never need direct file system access, and file-reading
    logic is centralised and reusable across all three agents.

EXPOSED TOOLS (3 total):
    1. read_document(file_path: str) -> str
       - Reads a local PDF, DOCX, or TXT file and returns its text content.
       - Validates file type and existence before reading.
       - Used by the Parser Agent to ingest the user's uploaded document.

    2. list_legal_references(category: str) -> str
       - Returns the content of a legal reference document from the local
         legal_library/ directory based on a category key.
       - Categories: "indemnification", "ip_ownership", "liability", "termination"
       - Used by the Risk Analyst Agent to compare clauses against standards.

    3. get_clause_template(clause_type: str) -> str
       - Retrieves a protective clause template from legal_library/templates/.
       - Used by the Protector Agent as a starting point for drafting
         safer alternative language.

TRANSPORT:
    This server uses the `stdio` transport — it is started via:
        python mcp_server/document_server.py
    The parent process (the agent runtime) communicates via stdin/stdout
    using the MCP wire protocol. FastMCP handles all protocol details.

RUNNING STANDALONE (for testing):
    python mcp_server/document_server.py

ERROR HANDLING:
    - All tools wrap their logic in try/except and return descriptive error
      strings (never raise exceptions) so agents can handle failures gracefully.
    - Unsupported file types return an explicit error message.
    - Missing files return a descriptive "not found" message.
"""

import os
import sys
import logging
from pathlib import Path

# FastMCP is the Python implementation of the Model Context Protocol.
# It provides a decorator-based API for defining tools that agents can call.
from mcp.server.fastmcp import FastMCP

# Document parsing libraries
import pdfplumber          # For reading PDF files
import docx                # For reading Microsoft Word .docx files

# Configure logging for the MCP server process.
# Since this runs as a subprocess, logs go to stderr (not stdout,
# which is reserved for MCP wire protocol messages).
logging.basicConfig(
    stream=sys.stderr,
    level=logging.INFO,
    format="[MCP-SERVER] %(asctime)s %(levelname)s: %(message)s"
)
logger = logging.getLogger(__name__)

# ==============================================================================
# CONFIGURATION
# ==============================================================================

# Base path of the project — used to resolve relative paths for the
# legal library regardless of what directory the server is started from.
_PROJECT_ROOT = Path(__file__).parent.parent.resolve()

# Legal library directory — contains reference standards and clause templates.
# Can be overridden by the LEGAL_LIBRARY_PATH environment variable.
_LEGAL_LIBRARY_PATH = Path(
    os.getenv("LEGAL_LIBRARY_PATH", str(_PROJECT_ROOT / "legal_library"))
)

# Mapping of category keys to reference document filenames.
# This allows agents to request references by logical name rather than
# knowing the exact file path structure.
_REFERENCE_MAP = {
    "indemnification": "indemnification_standards.md",
    "ip_ownership":    "ip_ownership_standards.md",
    "liability":       "liability_limits_standards.md",
    "termination":     "termination_clause_standards.md",
}

# Mapping of clause type keys to template filenames.
_TEMPLATE_MAP = {
    "indemnification":        "templates/indemnification_template.md",
    "ip":                     "templates/ip_template.md",
    "limitation_of_liability": "templates/limitation_of_liability_template.md",
}

# Maximum file size we'll read (10MB) — prevents memory issues with huge files.
_MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024


# ==============================================================================
# MCP SERVER INSTANTIATION
# ==============================================================================

# Create the FastMCP server instance.
# "LegalShieldDocumentServer" is the server name that appears in agent logs.
mcp = FastMCP("LegalShieldDocumentServer")

logger.info("LegalShield MCP Document Server initialising...")
logger.info("Legal library path: %s", _LEGAL_LIBRARY_PATH)


# ==============================================================================
# HELPER FUNCTIONS (Internal)
# ==============================================================================

def _read_pdf(file_path: Path) -> str:
    """
    Extracts text from a PDF file using pdfplumber.
    
    pdfplumber is preferred over PyPDF2 because it handles complex layouts,
    multi-column text, and tables much more accurately — critical for
    legal documents which often have structured, multi-column formatting.
    
    Args:
        file_path: Path to the PDF file.
    
    Returns:
        Extracted text as a single string with pages joined by newlines.
    """
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for page_num, page in enumerate(pdf.pages, 1):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(f"\n--- PAGE {page_num} ---\n{page_text}")
            else:
                logger.warning("Page %d of %s yielded no text (possibly image-based).", page_num, file_path.name)
    
    return "\n".join(text_parts)


def _read_docx(file_path: Path) -> str:
    """
    Extracts text from a Microsoft Word .docx file using python-docx.
    
    Reads both paragraph text and table cell text to ensure no content
    is missed — legal documents often use tables for clause definitions
    and schedules.
    
    Args:
        file_path: Path to the .docx file.
    
    Returns:
        Extracted text with paragraphs separated by newlines.
    """
    document = docx.Document(str(file_path))
    text_parts = []
    
    for para in document.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)
    
    # Also extract text from tables (common in legal documents)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text.strip())
    
    return "\n".join(text_parts)


def _read_txt(file_path: Path) -> str:
    """
    Reads a plain text (.txt) file.
    
    Attempts UTF-8 decoding first, then falls back to latin-1 which can
    handle any byte sequence (important for older legal documents with
    non-standard characters).
    
    Args:
        file_path: Path to the .txt file.
    
    Returns:
        The raw text content of the file.
    """
    try:
        return file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        logger.warning("UTF-8 decoding failed for %s; trying latin-1.", file_path.name)
        return file_path.read_text(encoding="latin-1")


# ==============================================================================
# MCP TOOL DEFINITIONS
# ==============================================================================

@mcp.tool()
def read_document(file_path: str) -> str:
    """
    Reads a local document file and returns its full text content.

    Supports PDF (.pdf), Microsoft Word (.docx), and plain text (.txt) files.
    Validates that the file exists, is accessible, and does not exceed the
    maximum allowed size (10MB) before reading.

    This is the PRIMARY tool used by the Parser Agent to ingest a contract
    document for analysis.

    Args:
        file_path: Absolute or relative path to the document file to read.

    Returns:
        The full text content of the document as a string.
        On error, returns a descriptive error message string (never raises).
    """
    logger.info("MCP tool 'read_document' called with path: %s", file_path)
    
    try:
        path = Path(file_path).resolve()
        
        # --- Validation ---
        if not path.exists():
            error_msg = f"ERROR: File not found: '{file_path}'. Please verify the path."
            logger.error(error_msg)
            return error_msg
        
        if not path.is_file():
            error_msg = f"ERROR: Path is not a file: '{file_path}'."
            logger.error(error_msg)
            return error_msg
        
        # Check file size before loading into memory
        file_size = path.stat().st_size
        if file_size > _MAX_FILE_SIZE_BYTES:
            error_msg = (
                f"ERROR: File too large ({file_size / 1024 / 1024:.1f} MB). "
                f"Maximum supported size is {_MAX_FILE_SIZE_BYTES / 1024 / 1024:.0f} MB."
            )
            logger.error(error_msg)
            return error_msg
        
        # --- File Type Detection & Reading ---
        suffix = path.suffix.lower()
        
        if suffix == ".pdf":
            logger.info("Reading PDF file: %s", path.name)
            text = _read_pdf(path)
        elif suffix == ".docx":
            logger.info("Reading DOCX file: %s", path.name)
            text = _read_docx(path)
        elif suffix == ".txt":
            logger.info("Reading TXT file: %s", path.name)
            text = _read_txt(path)
        else:
            error_msg = (
                f"ERROR: Unsupported file type '{suffix}'. "
                "LegalShield AI supports: .pdf, .docx, .txt"
            )
            logger.error(error_msg)
            return error_msg
        
        if not text or not text.strip():
            return (
                f"WARNING: The document '{path.name}' was read but contained no "
                "extractable text. It may be a scanned image PDF. "
                "Please provide a text-based PDF or DOCX file."
            )
        
        logger.info(
            "Successfully read '%s' — extracted %d characters.",
            path.name, len(text)
        )
        return text
    
    except Exception as e:
        error_msg = f"ERROR reading document '{file_path}': {type(e).__name__}: {e}"
        logger.exception(error_msg)
        return error_msg


@mcp.tool()
def list_legal_references(category: str) -> str:
    """
    Returns the content of a legal reference document for a given category.

    The legal library contains reference standards that define what constitutes
    a 'fair' or 'dangerous' clause in standard business contracts. The Risk
    Analyst Agent uses these standards to evaluate each extracted clause.

    Available categories:
    - "indemnification"  — Standards for indemnification clauses
    - "ip_ownership"     — Standards for intellectual property ownership
    - "liability"        — Standards for limitation of liability clauses
    - "termination"      — Standards for termination and exit clauses

    Args:
        category: One of the category keys listed above. Case-insensitive.

    Returns:
        The full text of the relevant legal reference document.
        If the category is unknown, returns a list of valid categories.
        On file error, returns a descriptive error message.
    """
    logger.info("MCP tool 'list_legal_references' called for category: '%s'", category)
    
    category_key = category.lower().strip()
    
    if category_key not in _REFERENCE_MAP:
        valid_cats = ", ".join(f'"{k}"' for k in _REFERENCE_MAP.keys())
        return (
            f"ERROR: Unknown category '{category}'. "
            f"Valid categories are: {valid_cats}"
        )
    
    reference_file = _LEGAL_LIBRARY_PATH / _REFERENCE_MAP[category_key]
    
    try:
        if not reference_file.exists():
            error_msg = (
                f"ERROR: Legal reference file not found: '{reference_file}'. "
                "Ensure the legal_library/ directory is intact."
            )
            logger.error(error_msg)
            return error_msg
        
        content = reference_file.read_text(encoding="utf-8")
        logger.info("Served legal reference for category '%s' (%d chars).", category_key, len(content))
        return content
    
    except Exception as e:
        error_msg = f"ERROR reading legal reference '{category}': {type(e).__name__}: {e}"
        logger.exception(error_msg)
        return error_msg


@mcp.tool()
def get_clause_template(clause_type: str) -> str:
    """
    Retrieves a protective clause template for a given clause type.

    These templates are professionally drafted protective clauses that the
    Protector Agent uses as a foundation when generating safer alternatives
    to risky clauses found in the user's contract.

    Available clause types:
    - "indemnification"         — Balanced mutual indemnification template
    - "ip"                      — Contractor-protective IP ownership template
    - "limitation_of_liability" — Standard liability cap template

    Args:
        clause_type: One of the clause type keys listed above. Case-insensitive.

    Returns:
        The full protective clause template text in Markdown format.
        If clause_type is unknown, returns a list of valid types.
        On file error, returns a descriptive error message.
    """
    logger.info("MCP tool 'get_clause_template' called for type: '%s'", clause_type)
    
    type_key = clause_type.lower().strip()
    
    if type_key not in _TEMPLATE_MAP:
        valid_types = ", ".join(f'"{k}"' for k in _TEMPLATE_MAP.keys())
        return (
            f"ERROR: Unknown clause type '{clause_type}'. "
            f"Valid types are: {valid_types}"
        )
    
    template_file = _LEGAL_LIBRARY_PATH / _TEMPLATE_MAP[type_key]
    
    try:
        if not template_file.exists():
            error_msg = (
                f"ERROR: Template file not found: '{template_file}'. "
                "Ensure the legal_library/templates/ directory is intact."
            )
            logger.error(error_msg)
            return error_msg
        
        content = template_file.read_text(encoding="utf-8")
        logger.info("Served clause template for type '%s' (%d chars).", type_key, len(content))
        return content
    
    except Exception as e:
        error_msg = f"ERROR reading clause template '{clause_type}': {type(e).__name__}: {e}"
        logger.exception(error_msg)
        return error_msg


# ==============================================================================
# SERVER ENTRY POINT
# ==============================================================================

if __name__ == "__main__":
    """
    Launches the MCP server when executed directly.
    
    The Google Antigravity SDK launches this script as a subprocess when
    an agent with McpStdioServer configured starts a session. The server
    communicates with the parent process via stdin/stdout (stdio transport).
    
    You can also run this directly for testing:
        python mcp_server/document_server.py
    """
    logger.info("Starting LegalShield MCP Document Server via stdio transport...")
    # mcp.run() starts the FastMCP server with stdio transport (default).
    # It blocks until the parent process closes the connection.
    mcp.run()
